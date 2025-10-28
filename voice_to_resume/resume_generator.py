import json
from docx import Document
from docx.shared import Pt

def generate_resume(json_file="user_data.json", output_file="generated_resume.docx"):
    # Load user data
    with open(json_file, "r", encoding="utf-8") as f:
        data = json.load(f)

    # Create a Word document
    doc = Document()
    doc.add_heading("Resume", level=0)

    for key, value in data.items():
        doc.add_heading(key, level=1)
        doc.add_paragraph(value)

    # Optional: set font size
    for para in doc.paragraphs:
        for run in para.runs:
            run.font.size = Pt(12)

    doc.save(output_file)
    print(f"💼 Resume generated successfully: {output_file}")
